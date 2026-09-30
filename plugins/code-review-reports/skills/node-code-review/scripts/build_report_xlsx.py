#!/usr/bin/env python3
"""Build the Core Review Report workbook (.xlsx) from a findings JSON file.

Usage:
    python build_report_xlsx.py findings.json
    python build_report_xlsx.py findings.json --out C:/path/REPORT.xlsx

The findings JSON shape is documented in ../references/excel-report.md.

The "Threshold Mapping" sheet is built per review, not shipped whole. The catalog
in threshold_mapping.json holds one block of thresholds and severity bands per
review area (architecture, security, performance, data, testing and so on). The
script keeps only the areas this review covers, worked out in this order:

    1. "threshold_areas" in the findings JSON, if present, by key, prefix, label
       or any of the area's "match" words, for example ["security", "PERF"].
    2. Otherwise, the Finding ID prefixes that actually appear in the rows, so a
       review with only SEC and ARCH findings gets only those two blocks.

"extra_thresholds" and "extra_severity_bands" append rows for a criterion the
catalog does not cover. "thresholds" and "severity_bands" replace the catalog
outright.

Requires openpyxl (pip install openpyxl).
"""

import argparse
import json
import logging
import math
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

LOGGER = logging.getLogger("build_report_xlsx")

HERE = os.path.dirname(os.path.abspath(__file__))
THRESHOLD_DATA_PATH = os.path.join(HERE, "threshold_mapping.json")

# --------------------------------------------------------------------------
# Fixed look of the workbook. Keep this in step with references/excel-report.md.
# --------------------------------------------------------------------------

HEADERS = [
    "Category",
    "Sub Category",
    "Module Name",
    "File / Location",
    "Type",
    "Metric",
    "Evidence",
    "Observation",
    "Severity",
    "Impact",
    "Recommendation",
    "Priority",
]

# Keys read from each row dict, in the same order as HEADERS.
ROW_KEYS = [
    "category",
    "sub_category",
    "module",
    "file",
    "type",
    "metric",
    "evidence",
    "observation",
    "severity",
    "impact",
    "recommendation",
    "priority",
]

COLUMN_WIDTHS = [32.1, 30.0, 24.6, 38.1, 18.3, 34.4, 82.4, 75.9, 13.4, 47.6, 81.7, 12.7]

# Columns that hold a short label. These are centred, the rest are left aligned.
CENTERED_COLUMNS = {1, 2, 3, 5, 9, 12}

HEADER_FILL = "FF1F3864"
HEADER_FONT_COLOR = "FFFFFFFF"
TITLE_FONT_COLOR = "FF1F3864"
SUBTITLE_FONT_COLOR = "FF595959"

BAND_FILLS = {
    "critical": "FFF4CCCC",  # light red
    "high": "FFFCE4D6",      # light orange
    "medium": "FFFFF2CC",    # light yellow
    "low": "FFE2EFDA",       # light green
}

SEVERITY_FILLS = {
    "Critical": BAND_FILLS["critical"],
    "High": BAND_FILLS["high"],
    "Medium": BAND_FILLS["medium"],
    "Low": BAND_FILLS["low"],
}

PRIORITY_FILLS = {
    "P1": BAND_FILLS["critical"],
    "P2": BAND_FILLS["medium"],
    "P3": BAND_FILLS["low"],
}

# Block 2 colours the four band columns from acceptable to critical.
BAND_COLUMN_FILLS = ["low", "medium", "high", "critical"]

VALID_TYPES = {"Route", "Middleware", "Function", "Module", "Config"}
VALID_SEVERITIES = ["Critical", "High", "Medium", "Low"]
VALID_PRIORITIES = ["P1", "P2", "P3"]
VALID_EFFORTS = {"Low", "Medium", "High"}

# Severity x Effort -> Priority, straight from references/rubrics.md.
PRIORITY_MATRIX = {
    ("Critical", "Low"): "P1",
    ("Critical", "Medium"): "P1",
    ("Critical", "High"): "P1",
    ("High", "Low"): "P1",
    ("High", "Medium"): "P1",
    ("High", "High"): "P2",
    ("Medium", "Low"): "P2",
    ("Medium", "Medium"): "P2",
    ("Medium", "High"): "P3",
    ("Low", "Low"): "P3",
    ("Low", "Medium"): "P3",
    ("Low", "High"): "P3",
}

FONT_NAME = "Calibri"
FONT_SIZE = 10
LINE_HEIGHT = 11.25
MIN_ROW_HEIGHT = 40.5
MAX_ROW_HEIGHT = 400.0
HEADER_ROW_HEIGHT = 33.75
MAX_SHEET_NAME = 31
ILLEGAL_SHEET_CHARS = set("[]:*?/\\")
THRESHOLD_WIDTHS = [32.0, 30.0, 38.0, 26.0, 34.0]
MAX_CELL_CHARS = 32767  # Excel refuses to open a file with a longer cell.

# Matches the Finding ID inside a File / Location cell, e.g. "`src/app.ts` (SEC-001)".
FINDING_ID_RE = re.compile(r"\(([A-Z]+)-\d+\)")

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------

def fill(color):
    """Return a solid PatternFill in the given ARGB colour."""
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def put_text(worksheet, row_index, column_index, value):
    """Write a value as plain text, so a leading '=' never becomes a formula."""
    text = str(value)
    if len(text) > MAX_CELL_CHARS:
        LOGGER.warning("Cell %s%d was %d chars and has been trimmed to fit Excel.",
                       get_column_letter(column_index), row_index, len(text))
        text = text[: MAX_CELL_CHARS - 3] + "..."
    cell = worksheet.cell(row=row_index, column=column_index, value=text)
    cell.data_type = "s"
    return cell


def wrapped_line_count(text, column_width):
    """Return roughly how many display lines a cell needs at that column width."""
    chars_per_line = max(8, int(column_width) - 2)
    total = 0
    for segment in str(text).split("\n"):
        total += max(1, math.ceil(len(segment) / chars_per_line))
    return total


def row_height_for(values, widths):
    """Return the row height in points that fits the tallest cell in the row."""
    lines = 1
    for value, width in zip(values, widths):
        if value in (None, ""):
            continue
        lines = max(lines, wrapped_line_count(value, width))
    return max(MIN_ROW_HEIGHT, min(MAX_ROW_HEIGHT, lines * LINE_HEIGHT))


def safe_sheet_name(name, used):
    """Return an Excel-legal, unique sheet name and record it in `used`."""
    cleaned = "".join(" " if ch in ILLEGAL_SHEET_CHARS else ch for ch in str(name)).strip()
    cleaned = cleaned[:MAX_SHEET_NAME] or "Sheet"
    candidate = cleaned
    suffix = 2
    while candidate.lower() in used:
        tail = " " + str(suffix)
        candidate = cleaned[: MAX_SHEET_NAME - len(tail)] + tail
        suffix += 1
    used.add(candidate.lower())
    return candidate


def load_threshold_catalog():
    """Load the per-area Threshold Mapping catalog shipped next to this script."""
    with open(THRESHOLD_DATA_PATH, "r", encoding="utf-8") as handle:
        return json.load(handle)


def area_lookup(catalog):
    """Return a dict of every name an area answers to, mapped to the area."""
    lookup = {}
    for area in catalog["areas"]:
        names = [area["key"], area["prefix"], area["label"]] + area.get("match", [])
        for name in names:
            lookup[name.strip().lower()] = area
    return lookup


def find_area(name, lookup):
    """Resolve one area name, allowing a criteria heading such as 'Security Review'.

    Tries an exact match first, then looks for a catalog name sitting inside the
    supplied text as a whole phrase. The longest matching name wins, so
    'api design review' picks API Design and not the shorter 'design'.
    """
    text = str(name).strip().lower()
    if text in lookup:
        return lookup[text]

    best = None
    best_length = 0
    for candidate, area in lookup.items():
        if len(candidate) <= 2:
            continue
        if re.search(r"(?<![a-z0-9])%s(?![a-z0-9])" % re.escape(candidate), text):
            if len(candidate) > best_length:
                best, best_length = area, len(candidate)
    return best


def prefixes_used(sheets):
    """Return the Finding ID prefixes that actually appear in the findings rows."""
    found = set()
    for sheet in sheets:
        for row in sheet.get("rows") or []:
            for prefix in FINDING_ID_RE.findall(str(row.get("file", ""))):
                found.add(prefix)
    return found


def resolve_areas(spec, catalog, sheets):
    """Pick the threshold areas for this review, in catalog order.

    `threshold_areas` in the spec wins. Otherwise the areas are inferred from the
    Finding ID prefixes on the findings sheets, so the sheet only ever carries
    thresholds that something was actually scored against.
    """
    lookup = area_lookup(catalog)
    wanted = spec.get("threshold_areas")

    if wanted:
        chosen = set()
        for name in wanted:
            area = find_area(name, lookup)
            if area is None:
                LOGGER.warning("Threshold area '%s' is not in the catalog and was skipped. "
                               "Add its rows with 'extra_thresholds' instead.", name)
            else:
                chosen.add(area["key"])
    else:
        chosen = set()
        for prefix in prefixes_used(sheets):
            area = lookup.get(prefix.lower())
            if area is None:
                LOGGER.warning("Finding ID prefix '%s' has no threshold area in the catalog. "
                               "Add its rows with 'extra_thresholds'.", prefix)
            else:
                chosen.add(area["key"])

    areas = [area for area in catalog["areas"] if area["key"] in chosen]
    if not areas:
        LOGGER.warning("No threshold area matched this review, falling back to every area.")
        areas = list(catalog["areas"])
    return areas


def build_threshold_tables(spec, catalog, sheets):
    """Return (intro, thresholds, band_intro, bands) for the Threshold Mapping sheet."""
    intro = spec.get("threshold_intro", catalog["threshold_intro"])
    band_intro = spec.get("band_intro", catalog["band_intro"])

    if "thresholds" in spec or "severity_bands" in spec:
        thresholds = spec.get("thresholds", [])
        bands = spec.get("severity_bands", [])
    else:
        areas = resolve_areas(spec, catalog, sheets)
        thresholds = [row for area in areas for row in area["thresholds"]]
        bands = [row for area in areas for row in area["severity_bands"]]
        labels = ", ".join(area["label"] for area in areas)
        intro = "%s Areas covered: %s." % (intro, labels)
        LOGGER.info("Threshold Mapping covers: %s", labels)

    thresholds = thresholds + spec.get("extra_thresholds", [])
    bands = bands + spec.get("extra_severity_bands", [])
    return intro, thresholds, band_intro, bands


def validate(sheets):
    """Return a list of readable problems. An empty list means the data is fine."""
    problems = []
    for sheet in sheets:
        sheet_name = sheet.get("name", "(unnamed sheet)")
        rows = sheet.get("rows") or []
        if not rows:
            problems.append("%s: has no rows" % sheet_name)
        for index, row in enumerate(rows, start=2):
            problems.extend(validate_row(row, "%s row %d" % (sheet_name, index)))
    return problems


def validate_row(row, where):
    """Return the problems found in one findings row, tagged with its location."""
    problems = []
    for key in ROW_KEYS:
        if not str(row.get(key, "")).strip():
            problems.append("%s: '%s' is empty" % (where, key))

    row_type = row.get("type")
    if row_type and row_type not in VALID_TYPES:
        problems.append("%s: type '%s' is not one of %s"
                        % (where, row_type, ", ".join(sorted(VALID_TYPES))))

    severity = row.get("severity")
    if severity and severity not in VALID_SEVERITIES:
        problems.append("%s: severity '%s' is not one of %s"
                        % (where, severity, ", ".join(VALID_SEVERITIES)))

    priority = row.get("priority")
    if priority and priority not in VALID_PRIORITIES:
        problems.append("%s: priority '%s' is not one of %s"
                        % (where, priority, ", ".join(VALID_PRIORITIES)))

    effort = row.get("effort")
    if not effort:
        return problems
    if effort not in VALID_EFFORTS:
        problems.append("%s: effort '%s' is not Low, Medium or High" % (where, effort))
    elif severity in VALID_SEVERITIES and priority in VALID_PRIORITIES:
        expected = PRIORITY_MATRIX[(severity, effort)]
        if expected != priority:
            problems.append("%s: priority %s does not match Severity %s x Effort %s, expected %s"
                            % (where, priority, severity, effort, expected))
    return problems


# --------------------------------------------------------------------------
# Sheet writers
# --------------------------------------------------------------------------

def write_findings_sheet(worksheet, rows, evidence_monospace):
    """Write one service's findings as the 12-column table, styled and filtered."""
    header_font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=HEADER_FONT_COLOR)
    header_fill = fill(HEADER_FILL)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column_index, title in enumerate(HEADERS, start=1):
        cell = put_text(worksheet, 1, column_index, title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = BORDER
    worksheet.row_dimensions[1].height = HEADER_ROW_HEIGHT

    for column_index, width in enumerate(COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    body_font = Font(name=FONT_NAME, size=FONT_SIZE)
    code_font = Font(name="Consolas", size=9)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(vertical="center", wrap_text=True)

    for offset, row in enumerate(rows):
        excel_row = offset + 2
        values = [str(row.get(key, "")) for key in ROW_KEYS]
        for column_index, value in enumerate(values, start=1):
            cell = put_text(worksheet, excel_row, column_index, value)
            is_evidence = column_index == 7 and evidence_monospace
            cell.font = code_font if is_evidence else body_font
            cell.alignment = centered if column_index in CENTERED_COLUMNS else left
            cell.border = BORDER
        worksheet.cell(row=excel_row, column=9).fill = fill(SEVERITY_FILLS[row["severity"]])
        worksheet.cell(row=excel_row, column=12).fill = fill(PRIORITY_FILLS[row["priority"]])
        worksheet.row_dimensions[excel_row].height = row_height_for(values, COLUMN_WIDTHS)

    worksheet.auto_filter.ref = "A1:L%d" % (len(rows) + 1)
    worksheet.freeze_panes = "F2"


def write_block_title(worksheet, row_index, text, font):
    """Write a block title or subtitle into column A of the threshold sheet."""
    cell = put_text(worksheet, row_index, 1, text)
    cell.font = font
    cell.alignment = Alignment(vertical="center")


def write_block_header(worksheet, row_index, titles):
    """Write a dark header row for one block of the threshold sheet."""
    font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color=HEADER_FONT_COLOR)
    for column_index, title in enumerate(titles, start=1):
        cell = put_text(worksheet, row_index, column_index, title)
        cell.font = font
        cell.fill = fill(HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER


def write_threshold_rows(worksheet, start_row, thresholds):
    """Write the block 1 threshold rows and return the next free row index."""
    body_font = Font(name=FONT_NAME, size=FONT_SIZE)
    centered = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(vertical="center", wrap_text=True)

    row_index = start_row
    for threshold_row in thresholds:
        values = [str(value) for value in threshold_row]
        for column_index, value in enumerate(values, start=1):
            cell = put_text(worksheet, row_index, column_index, value)
            cell.font = body_font
            cell.alignment = centered if column_index in (4, 5) else left
            cell.border = BORDER
        severity, priority = values[3], values[4]
        if severity in SEVERITY_FILLS:
            worksheet.cell(row=row_index, column=4).fill = fill(SEVERITY_FILLS[severity])
        if priority in PRIORITY_FILLS:
            worksheet.cell(row=row_index, column=5).fill = fill(PRIORITY_FILLS[priority])
        worksheet.row_dimensions[row_index].height = row_height_for(values, THRESHOLD_WIDTHS)
        row_index += 1
    return row_index


def write_band_rows(worksheet, start_row, bands):
    """Write the block 2 severity band rows and return the next free row index."""
    body_font = Font(name=FONT_NAME, size=FONT_SIZE)
    left = Alignment(vertical="center", wrap_text=True)

    row_index = start_row
    for band_row in bands:
        values = [str(value) for value in band_row]
        for column_index, value in enumerate(values, start=1):
            cell = put_text(worksheet, row_index, column_index, value)
            cell.font = body_font
            cell.alignment = left
            cell.border = BORDER
            if column_index >= 2:
                cell.fill = fill(BAND_FILLS[BAND_COLUMN_FILLS[column_index - 2]])
        worksheet.row_dimensions[row_index].height = row_height_for(values, THRESHOLD_WIDTHS)
        row_index += 1
    return row_index


def write_threshold_sheet(worksheet, tables):
    """Write the two-block Threshold Mapping sheet that explains the ratings."""
    threshold_intro, thresholds, band_intro, bands = tables

    for column_index, width in enumerate(THRESHOLD_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    title_font = Font(name=FONT_NAME, size=11, bold=True, color=TITLE_FONT_COLOR)
    subtitle_font = Font(name=FONT_NAME, size=9, color=SUBTITLE_FONT_COLOR)

    row_index = 1
    write_block_title(worksheet, row_index, "BLOCK 1  |  AUDIT METRIC THRESHOLDS AND TREATMENT",
                      title_font)
    write_block_title(worksheet, row_index + 1, threshold_intro, subtitle_font)
    write_block_header(worksheet, row_index + 2,
                       ["Audit Metric", "Recommended Threshold", "Audit Treatment",
                        "Severity", "Priority"])
    row_index = write_threshold_rows(worksheet, row_index + 3, thresholds)

    row_index += 2
    write_block_title(worksheet, row_index, "BLOCK 2  |  SEVERITY BAND TABLE", title_font)
    write_block_title(worksheet, row_index + 1, band_intro, subtitle_font)
    write_block_header(worksheet, row_index + 2,
                       ["Metric", "Low / Acceptable", "Medium", "High", "Critical / Very High"])
    write_band_rows(worksheet, row_index + 3, bands)

    worksheet.freeze_panes = "A4"


# --------------------------------------------------------------------------
# Entry point
# --------------------------------------------------------------------------

def build(spec, output_override=None):
    """Validate the spec, write the workbook, and return (output path, sheets)."""
    sheets = spec.get("sheets") or []
    if not sheets:
        raise SystemExit("ERROR: the findings JSON has no 'sheets'.")

    problems = validate(sheets)
    if problems:
        LOGGER.error("Findings data did not pass validation, nothing was written:")
        for problem in problems:
            LOGGER.error("  - %s", problem)
        raise SystemExit(1)

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names = set()
    evidence_monospace = bool(spec.get("evidence_monospace", False))

    for sheet in sheets:
        name = safe_sheet_name(sheet.get("name", "Findings"), used_names)
        write_findings_sheet(workbook.create_sheet(name), sheet["rows"], evidence_monospace)

    if spec.get("include_threshold_sheet", True):
        name = safe_sheet_name("Threshold Mapping", used_names)
        catalog = load_threshold_catalog()
        tables = build_threshold_tables(spec, catalog, sheets)
        write_threshold_sheet(workbook.create_sheet(name), tables)

    output = os.path.abspath(output_override or spec.get("output") or "CORE_REVIEW_REPORT.xlsx")
    directory = os.path.dirname(output)
    if directory and not os.path.isdir(directory):
        os.makedirs(directory)
    workbook.save(output)
    return output, sheets


def summarize(output, sheets):
    """Log the output path and a per-sheet row and severity count."""
    LOGGER.info("Wrote %s", output)
    for sheet in sheets:
        counts = {}
        for row in sheet["rows"]:
            counts[row["severity"]] = counts.get(row["severity"], 0) + 1
        summary = ", ".join("%s %d" % (severity, counts[severity])
                            for severity in VALID_SEVERITIES if severity in counts)
        LOGGER.info("  %-32s %3d rows  (%s)",
                    sheet["name"], len(sheet["rows"]), summary or "no findings")


def main():
    """Read the findings JSON named on the command line and build the workbook."""
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    parser = argparse.ArgumentParser(description="Build the Core Review Report workbook.")
    parser.add_argument("json_path", help="Path to the findings JSON file.")
    parser.add_argument("--out", dest="out", default=None,
                        help="Override the output path from the JSON.")
    args = parser.parse_args()

    with open(args.json_path, "r", encoding="utf-8") as handle:
        spec = json.load(handle)

    summarize(*build(spec, args.out))


if __name__ == "__main__":
    main()
